"""Unit Test Suite for Multimodal Document Parser with Column Layout Reconstruction,

Cross-Page Table Merging, and Multi-Sheet Multi-Level Excel Headers.
"""

from unittest.mock import patch

import openpyxl
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from backend.core.multimodal_parser import MultimodalDocumentParser, ParsedChunk
from backend.tasks.celery_app import celery_app
from backend.tasks.ingestion_tasks import async_parse_multimodal_file

# Configure Celery Task Eager mode for tests
celery_app.conf.update(
    task_always_eager=True,
    task_eager_propagates=True,
    task_store_eager_result=True,
    result_backend="cache+memory://",
)


def _create_sample_multicol_pdf(file_path):
    """Generates a synthetic PDF with two distinct text columns and markdown table syntax."""
    c = canvas.Canvas(str(file_path), pagesize=letter)
    # Page 1: multi-column format with large spaces between columns
    c.drawString(50, 750, "LeftCol Paragraph 1 Content        RightCol Heading Title")
    c.drawString(50, 730, "LeftCol Paragraph 2 Additional     RightCol Description Body")
    c.drawString(50, 710, "LeftCol Paragraph 3 Conclusion     RightCol Footer Note")

    # Add a markdown-styled table with ascii headers
    c.drawString(50, 650, "| Department | Budget_Limit | Approver |")
    c.drawString(50, 630, "| Engineering | 20000 | Zhang_Director |")
    c.drawString(50, 610, "| Marketing | 15000 | Li_Manager |")

    c.showPage()
    # Page 2: Continuation table
    c.drawString(50, 750, "| Department | Budget_Limit | Approver |")
    c.drawString(50, 730, "| Operations | 10000 | Wang_Lead |")
    c.drawString(50, 710, "| Administration | 5000 | Zhao_Lead |")
    c.showPage()
    c.save()


def _create_sample_hierarchical_excel(file_path):
    """Generates a synthetic Excel workbook with multi-sheet and merged header cells."""
    wb = openpyxl.Workbook()

    # Sheet 1: 部门采购与报销 (Multi-level header)
    ws1 = wb.active
    ws1.title = "采购报销标准"

    # Merge top row for categories
    ws1.merge_cells("A1:B1")
    ws1["A1"] = "基本信息"
    ws1.merge_cells("C1:D1")
    ws1["C1"] = "财务预算控制"

    # Row 2 sub-headers
    ws1["A2"] = "部门名称"
    ws1["B2"] = "岗位职级"
    ws1["C2"] = "单笔限额"
    ws1["D2"] = "年度总预算"

    # Row 3 & 4 data
    ws1.append(["研发部", "资深架构师", "20000", "500000"])
    ws1.append(["销售部", "大区总监", "30000", "800000"])

    # Sheet 2: 资产清单 (Second Sheet)
    ws2 = wb.create_sheet(title="IT资产清单")
    ws2.append(["资产编号", "设备型号", "领用人", "状态"])
    ws2.append(["AST-001", "MacBook Pro M3", "张三", "在用"])
    ws2.append(["AST-002", "Dell UltraSharp 4K", "李四", "闲置"])

    wb.save(str(file_path))


class TestMultimodalDocumentParser:
    """Test suite verifying PDF column reconstruction, cross-page tables, and Excel parsing."""

    def test_parsed_chunk_model(self):
        chunk = ParsedChunk(
            content="### 部门规章说明",
            chunk_type="text",
            metadata={"source": "test.pdf", "page": 1},
        )
        assert chunk.content == "### 部门规章说明"
        assert chunk.chunk_type == "text"
        assert chunk.metadata["page"] == 1

        langchain_doc = chunk.to_langchain_document()
        assert langchain_doc.page_content == "### 部门规章说明"
        assert langchain_doc.metadata["chunk_type"] == "text"
        assert langchain_doc.metadata["source"] == "test.pdf"

    def test_pdf_multicol_and_crosspage_table_parsing(self, tmp_path):
        pdf_path = tmp_path / "enterprise_policy.pdf"
        _create_sample_multicol_pdf(pdf_path)

        chunks = MultimodalDocumentParser.parse_pdf(
            pdf_path,
            detect_columns=True,
            merge_cross_page_tables=True,
        )

        assert len(chunks) >= 1
        # Verify table chunk extraction
        table_chunks = [c for c in chunks if c.chunk_type == "table"]
        assert len(table_chunks) >= 1

        first_table = table_chunks[0]
        assert "Department" in first_table.content
        assert "Budget_Limit" in first_table.content
        assert "Engineering" in first_table.content
        assert "table_meta" in first_table.metadata
        assert first_table.metadata["table_meta"]["columns"] == ["Department", "Budget_Limit", "Approver"]
        # Verify cross-page table merge
        assert first_table.metadata["table_meta"]["is_merged_cross_page"] is True
        assert first_table.metadata["table_meta"]["row_count"] == 4

    def test_excel_multisheet_and_hierarchical_headers(self, tmp_path):
        excel_path = tmp_path / "budget_and_assets.xlsx"
        _create_sample_hierarchical_excel(excel_path)

        chunks = MultimodalDocumentParser.parse_excel(excel_path, max_header_rows=2)

        assert len(chunks) == 2
        # Sheet 1 Check
        sheet1_chunk = next(c for c in chunks if c.metadata.get("sheet_name") == "采购报销标准")
        assert sheet1_chunk.chunk_type == "table"
        assert "基本信息_部门名称" in sheet1_chunk.content or "部门名称" in sheet1_chunk.content
        assert "财务预算控制_单笔限额" in sheet1_chunk.content or "单笔限额" in sheet1_chunk.content
        assert "研发部" in sheet1_chunk.content
        assert "资深架构师" in sheet1_chunk.content

        # Sheet 2 Check
        sheet2_chunk = next(c for c in chunks if c.metadata.get("sheet_name") == "IT资产清单")
        assert sheet2_chunk.chunk_type == "table"
        assert "MacBook Pro M3" in sheet2_chunk.content
        assert "AST-001" in sheet2_chunk.content

    def test_universal_parse_document_dispatch(self, tmp_path):
        # Markdown test
        md_file = tmp_path / "guide.md"
        md_file.write_text(
            "# 员工指南\n\n| 事项 | 要求 |\n|---|---|\n| 考勤 | 9:00前打卡 |\n\n这是指南详细正文内容描述段落，用于验证段落提取。",
            encoding="utf-8",
        )
        chunks = MultimodalDocumentParser.parse_document(md_file)
        assert len(chunks) >= 2
        types = [c.chunk_type for c in chunks]
        assert "table" in types
        assert "text" in types

    @patch("backend.tasks.ingestion_tasks.upsert_documents")
    def test_async_parse_multimodal_file_task(self, mock_upsert, tmp_path):
        mock_upsert.return_value = 2
        excel_path = tmp_path / "test_async_assets.xlsx"
        _create_sample_hierarchical_excel(excel_path)

        task = async_parse_multimodal_file.delay(
            file_path_str=str(excel_path),
            detect_columns=True,
            merge_cross_page_tables=True,
            index_to_vector_store=True,
            category="financial_budget",
        )

        assert task.status == "SUCCESS"
        result = task.result
        assert result["status"] == "COMPLETED"
        assert result["filename"] == "test_async_assets.xlsx"
        assert result["chunks_count"] == 2
        assert result["table_chunks_count"] == 2
        assert result["chunks_indexed"] == 2
