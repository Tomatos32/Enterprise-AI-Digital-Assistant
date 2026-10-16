"""RAG Document Loader & Parser Utilities for Enterprise AI Digital Assistant."""

from pathlib import Path
from typing import List, Dict, Any
import openpyxl
from pypdf import PdfReader
from langchain_core.documents import Document


def extract_metadata_from_filename(file_path: Path) -> Dict[str, Any]:
    """Extracts default metadata fields based on document filename conventions."""
    name = file_path.name.upper()
    category = "GENERAL"
    if "SECURITY" in name or "安全" in name or "保密" in name:
        category = "SECURITY"
    elif "CS" in name or "客服" in name or "售后" in name:
        category = "CS"
    elif "HR" in name or "人事" in name or "行政" in name:
        category = "HR"
    elif "PROCUREMENT" in name or "PROC" in name or "采购" in name:
        category = "PROCUREMENT"
    elif "FINANCE" in name or "财务" in name or "报销" in name:
        category = "FINANCE"
    elif "IT_" in name or "_IT" in name or "IT" in name.split("_") or "运维" in name:
        category = "IT"

    return {
        "source": str(file_path),
        "filename": file_path.name,
        "category": category,
        "format": file_path.suffix.lstrip(".").lower()
    }


def load_markdown_document(file_path: Path) -> List[Document]:
    """Loads a Markdown file into Document objects."""
    text = file_path.read_text(encoding="utf-8")
    meta = extract_metadata_from_filename(file_path)
    meta["title"] = text.splitlines()[0].lstrip("#").strip() if text else file_path.stem
    return [Document(page_content=text, metadata=meta)]


def load_pdf_document(file_path: Path) -> List[Document]:
    """Loads a PDF file into Document objects."""
    reader = PdfReader(str(file_path))
    meta = extract_metadata_from_filename(file_path)
    documents = []
    
    for idx, page in enumerate(reader.pages):
        page_text = page.extract_text() or ""
        page_meta = meta.copy()
        page_meta["page"] = idx + 1
        documents.append(Document(page_content=page_text, metadata=page_meta))
        
    return documents


def load_excel_document(file_path: Path) -> List[Document]:
    """Loads an Excel file and converts sheet data into Markdown table Document."""
    wb = openpyxl.load_workbook(str(file_path), data_only=True)
    meta = extract_metadata_from_filename(file_path)
    documents = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        
        # Build Markdown Table representation
        header = rows[0]
        header_line = "| " + " | ".join(str(cell or "") for cell in header) + " |"
        sep_line = "| " + " | ".join("---" for _ in header) + " |"
        
        body_lines = []
        for r in rows[1:]:
            row_str = "| " + " | ".join(str(cell or "") for cell in r) + " |"
            body_lines.append(row_str)
            
        md_table_text = f"# Excel Sheet: {sheet_name}\n\n" + "\n".join([header_line, sep_line] + body_lines)
        sheet_meta = meta.copy()
        sheet_meta["sheet"] = sheet_name
        documents.append(Document(page_content=md_table_text, metadata=sheet_meta))

    return documents


def load_sample_kb_documents(kb_dir: Path) -> List[Document]:
    """Recursively loads all sample_kb files (Markdown, PDF, Excel)."""
    docs: List[Document] = []
    if not kb_dir.exists():
        return docs

    for path in kb_dir.rglob("*"):
        if path.is_file():
            suffix = path.suffix.lower()
            if suffix == ".md":
                docs.extend(load_markdown_document(path))
            elif suffix == ".pdf":
                docs.extend(load_pdf_document(path))
            elif suffix in [".xlsx", ".xls"]:
                docs.extend(load_excel_document(path))

    return docs
